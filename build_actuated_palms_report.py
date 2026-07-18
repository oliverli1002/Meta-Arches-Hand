from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKDIR = Path(r"C:\Users\Oliver.OLIVER-LI-DESKT\Desktop\变胞手论文绘图")
SOURCE_PDF = WORKDIR / "人手掌弓生理结构文章" / "2024-tmech-Actuated_Palms_for_Soft_Robotic_Hands_Review_and_Perspectives.pdf"
OUTPUT_DOCX = WORKDIR / "Actuated_Palms_for_Soft_Robotic_Hands_中文研读报告.docx"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x0B, 0x25, 0x45)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"


def set_font(run, latin="Calibri", east_asia="SimSun"):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def style_run(run, size=None, bold=None, italic=None, color=None):
    set_font(run)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def para_tokens(paragraph, before=0, after=6, line=1.25):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_para(doc, text="", style=None, before=0, after=6, line=1.25, bold=False, italic=False, color=None, align=None):
    p = doc.add_paragraph(style=style)
    para_tokens(p, before, after, line)
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        style_run(run, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_font(run)
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_in, indent_dxa=120):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    widths_dxa = [int(round(w * 1440)) for w in widths_in]
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[idx])
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths_in, font_size=8.8, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_in)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, header_fill)
        p = cell.paragraphs[0]
        para_tokens(p, after=0, line=1.12)
        r = p.add_run(header)
        style_run(r, size=font_size, bold=True, color=INK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            para_tokens(p, after=0, line=1.12)
            r = p.add_run(str(value))
            style_run(r, size=font_size)
    set_table_geometry(table, widths_in)
    add_para(doc, "", after=2)
    return table


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "Actuated Palms for Soft Robotic Hands - 中文研读报告"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para_tokens(header, after=0, line=1.0)
    for run in header.runs:
        style_run(run, size=9, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.text = "Academic Paper Analysis | 2026-06-28"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para_tokens(footer, after=0, line=1.0)
    for run in footer.runs:
        style_run(run, size=9, color=GRAY)


def add_cover(doc):
    add_para(doc, "中文研读报告", before=10, after=4, bold=True, color=BLUE)
    title = doc.add_paragraph()
    para_tokens(title, after=6, line=1.15)
    r = title.add_run("Actuated Palms for Soft Robotic Hands: Review and Perspectives")
    style_run(r, size=21, bold=True, color=INK)
    add_para(doc, "主动掌软体机器人手：综述与展望", after=14, italic=True, color=GRAY)

    rows = [
        ("作者", "Maria Pozzi; Monica Malvezzi; Domenico Prattichizzo; Gionata Salvietti"),
        ("期刊", "IEEE/ASME Transactions on Mechatronics, Vol. 29, No. 2, pp. 902-912"),
        ("年份/DOI", "2024; DOI: 10.1109/TMECH.2023.3328944"),
        ("源 PDF", str(SOURCE_PDF)),
        ("文献性质", "综述论文。作者系统梳理带主动掌的软体多指机器人手/夹爪，并按手指结构、拟人程度、驱动方式、掌自由度和功能进行分类。"),
        ("提取说明", "PDF 共 11 页，文本可抽取；少量连字在文本层显示异常，报告中已按上下文恢复为 finger, flexion, configuration 等术语。"),
        ("与你课题的关联", "这篇文章直接支撑“机器人手掌不应只是刚性基座，而可作为主动重构部件参与包络抓持和手内操作”的设计动机。"),
    ]
    add_table(doc, ["项目", "内容"], rows, [1.35, 5.15], font_size=8.8, header_fill=LIGHT_GRAY)

    add_heading(doc, "阅读导引", 1)
    add_para(
        doc,
        "这篇 T-MECH 综述的核心观点非常贴近你的变胞手论文：软体/欠驱动机器人手如果只增加手指自由度，结构、驱动和控制会迅速复杂化；而在手掌中加入少量主动自由度，往往能更高效地提升包络抓持、拇指对掌、手指基座重构和手内操作能力。",
    )
    add_para(
        doc,
        "对你的论文最关键的不是它列出的所有样机，而是它的论证路径：先说明人手掌在抓持与操作中的主动作用，再说明机器人手长期把掌部当作被动支架的不足，最后用综述证据证明“主动掌”已成为软体/灵巧手设计的重要趋势。这条路径可以直接服务于你的 Introduction 和 Related Work。",
    )
    doc.add_page_break()


def add_metadata_and_terms(doc):
    add_heading(doc, "文献信息", 1)
    rows = [
        ("研究主题", "带主动掌的软体机器人手/夹爪的设计综述与未来趋势。"),
        ("研究对象", "软体多指机器人手，包括 articulated soft hands 和 continuous soft hands；进一步区分拟人型与非拟人型设计。"),
        ("检索方法", "在 Scopus 中以 robotic hand AND palm 等关键词检索；初筛 150 篇，按标准保留 17 篇；另检索 2018-2022 年机器人手/夹爪综述 53 篇，补充 2 篇；再通过引文追踪，最终纳入 39 篇。"),
        ("分类维度", "手类型、拟人程度、掌驱动方式、电/腱绳/气动、掌驱动自由度 DoA、掌主要功能、掌材料刚柔属性。"),
        ("核心结论", "主动掌通常只需 1-2 个驱动自由度，却能显著提升软体手的抓取稳定性、包络能力、手内操作和可重构抓取范围。"),
    ]
    add_table(doc, ["条目", "内容"], rows, [1.35, 5.15], font_size=8.8)

    add_heading(doc, "术语表", 1)
    terms = [
        ("actuated palm", "主动掌/驱动掌", "掌部含有主动自由度，并直接参与抓取或操作，而非仅作为手指安装基座。"),
        ("soft robotic hand", "软体机器人手", "含有柔顺结构或软材料的多指机器人手。文中包括关节柔顺型和连续变形型。"),
        ("articulated soft hand", "关节型软体手", "柔顺性主要位于关节，手指仍由较明确的关节/连杆结构构成。"),
        ("continuous soft hand", "连续型软体手", "手指或掌部呈连续可变形结构，常由气动腔、硅胶、柔性材料等构成。"),
        ("DoA", "驱动自由度", "Degree of Actuation，指掌部主动驱动输入数量，不等同于全部运动自由度。"),
        ("finger reconfiguration", "手指重构", "通过掌部机构改变手指基座位置或姿态，从而改变工作空间和抓取类型。"),
        ("enveloping grasp", "包络抓持", "通过更大的接触斑块或更多接触区域包围物体，提高抓持鲁棒性。"),
        ("thumb opposition", "拇指对掌", "拇指与其他手指/掌面形成相对关系，是拟人手抓持的关键功能。"),
        ("palm hollowing", "掌心凹陷/掌部凹化", "掌部主动弯曲，使环指和小指侧向拇指方向靠近，类似人手掌弓形成凹腔。"),
        ("in-hand manipulation", "手内操作", "物体已被抓持后，在手内滑动、翻转、旋转或重新定位。"),
    ]
    add_table(doc, ["英文术语", "中文译名", "说明"], terms, [1.45, 1.35, 3.70], font_size=8.4)


def add_section_reading(doc):
    add_heading(doc, "中文结构化译读", 1)
    add_para(doc, "说明：以下按原论文结构进行中文译读与压缩翻译，保留章节逻辑、关键术语、图表含义和对你课题有用的论据。", italic=True, color=GRAY)

    add_heading(doc, "摘要", 2)
    add_para(
        doc,
        "作者指出，人手掌在抓取和操作任务中具有关键且主动的作用，但许多机器人手仍把掌部设计为被动支架，只用于安装手指。既有研究更多关注灵巧手指的设计，而较少关注掌部自身的运动。软体技术为多指机器人夹爪的设计和驱动提供了新机会，越来越多装置开始具备主动掌，即掌部包含主动自由度并在抓取任务中发挥作用。因此，本文对具有主动掌的软体机器人手样机进行分类和分析，并讨论未来设计方向。",
    )

    add_heading(doc, "I. Introduction", 2)
    add_para(
        doc,
        "多指机器人手的机电设计涉及运动学结构、材料、传感和驱动系统等多个挑战。不同于多数综述聚焦整个机器人手，本文聚焦一个更具体却常被忽略的部分：手掌。虽然已有大量工作研究具有不同形状和功能的机器人手指，但真正把机器人掌部设计和驱动作为研究对象的工作相对较少。传统设计中，手掌通常是固定部件，支撑手指并容纳驱动系统；即使拟人机器人手实现了一定程度的拇指对掌，也往往依赖拇指根部关节而非真正移动掌部。",
    )
    add_para(
        doc,
        "作者强调，在机器人掌中加入主动自由度可以增加或增强有用功能，尤其适用于手指欠驱动、灵巧性受限的软体手。文章综述对象被限定为软体多指机器人手中的主动掌，即至少具有两个手指、包含专门设计的柔顺元件，并且掌部主动自由度在抓取或操作中发挥作用的装置。图 1 展示了几类典型用途：新抓取策略、包络物体、手内操作和拇指对掌。",
    )

    add_heading(doc, "II. Palm of the Human Hand", 2)
    add_para(
        doc,
        "这一节对你的论文最重要。作者先说明人手是动作和感知的双重界面，掌部参与二者：掌侧皮肤具备较高触觉神经密度，掌内肌肉、掌骨和关节共同决定手掌运动。人手掌结构启发了拟人和非拟人机器人夹爪设计，使主动控制掌部额外自由度成为可能，从而包络物体或以可控方式操作物体。",
    )
    add_para(
        doc,
        "作者将掌部运动的生理解剖基础总结为三条主要掌弓：纵弓、远端横弓和斜弓。纵弓从腕横纹延伸至中指或示指指尖；远端横弓在示指、中指、环指和小指掌骨头处形成横向凹曲；斜弓由可对掌拇指与其他手指共同形成。肌肉使掌部能沿这些弓形方向弯曲并形成凹形掌面，这对于执握任务中包绕物体至关重要。",
    )
    add_para(
        doc,
        "作者还强调，掌部不仅因为可运动而重要，其组织构型和柔软一致性也使其适于抓取和包络物体。掌侧皮肤较厚，具有软垫状皮下脂肪组织，且与掌腱膜连接较牢固；这让掌面在大接触面积任务中能够作为对向支撑面。人类抓取分类也表明，在 power-palm 和 power-pad 等需要大接触斑块的抓取中，掌部作为重要对向面；在手内操作中，掌部结构对于建立大面积接触同样关键。",
    )
    add_para(
        doc,
        "这段可直接支撑你的论断：人手掌不是刚性平面，而是通过远端横弓、纵弓和斜弓形成可调节的凹腔；这一非平面掌面在包络物体、扩大接触区域和手内操作中有功能意义。",
    )

    add_heading(doc, "III. Methodology", 2)
    add_para(
        doc,
        "作者采用 Scopus 检索和引文追踪相结合的方法。第一条检索路径围绕 robotic hand 与 palm，初步得到 150 条结果，其中 17 篇符合综述标准。第二条路径检索 2018-2022 年机器人夹爪/手综述，得到 53 条结果并补充 2 篇相关文献。随后检查这些文献的引用和被引关系，最终纳入 39 篇。分类框架基于软体手类型：关节型 soft hands 与连续型 soft hands；每类又分为非拟人型 grippers 和拟人型 hands。作者还记录了掌驱动方式、掌 DoA、功能、拟人程度和材料类型。",
    )

    add_heading(doc, "IV. Actuated Palms for Articulated Soft Hands", 2)
    add_heading(doc, "A. 非拟人关节型软体夹爪", 3)
    add_para(
        doc,
        "这一类主动掌多用于改变手指基座的相对位置或方向。典型方式包括齿轮机构、五杆机构、腱驱动结构或柔性铰链，使侧指旋转、平移或改变配置，从而实现球形抓取、圆柱抓取、两指平行夹取等不同抓取类型。另一些主动掌并不追求拟人结构，而是加入可动 scoop 结构，利用环境约束实现新的抓取策略。",
    )
    add_heading(doc, "B. 拟人关节型软体手", 3)
    add_para(
        doc,
        "拟人关节型软体手的主动掌更接近人手功能，常通过拇指底部或掌部结构辅助拇指对掌。例如 Pisa/IIT SoftHand 系列采用柔性掌结构来改善物体包络。相关工作显示，在人体手掌启发的人工皮肤材料、柔性掌和欠驱动手指协同作用下，手可以用有限驱动实现较强适应性和稳定抓持。",
    )

    add_heading(doc, "V. Actuated Palms for Continuous Soft Hands", 2)
    add_heading(doc, "A. 非拟人连续型软体夹爪", 3)
    add_para(
        doc,
        "连续型软体夹爪中的主动掌常由气动腔、硅胶楔形结构、可伸缩机构或可变摩擦/预紧结构构成。其作用包括改变夹爪工作空间、改变掌部直径和高度、支撑不同尺寸物体、调节掌面与物体之间的摩擦系数或法向预载。综述指出，某些真空驱动可弯掌能增加夹爪工作空间、负载能力和接触面积；可变摩擦掌则能让相同手指运动在不同掌驱动状态下产生滑动或翻转等不同手内操作。",
    )
    add_heading(doc, "B. 拟人连续型软体手", 3)
    add_para(
        doc,
        "拟人连续型软体手中，RBO Hand 2 和 RBO Hand 3 是重要案例。RBO Hand 2 具有两个气动掌部执行器，刚度高于手指，可支撑被抓物并产生不同拇指对掌运动；掌部运动也帮助整手更好地包绕物体。RBO Hand 3 进一步采用生物启发设计，掌部具有主动自由度，可实现 palm hollowing，即掌部屈曲并使环指、小指向拇指方向靠近。作者指出，这一掌部驱动对通过 Kapandji test 很关键。",
    )
    add_para(
        doc,
        "另一些软体拟人手采用单体式气动掌结构，直接在掌中布置气腔以复现人手掌弓运动。例如 Yang 等实现了纵弓和远端横弓；Firth 等聚焦纵弓和斜弓。也有工作通过颗粒堵塞或多层软掌结构使掌面更好地稳定被抓物。对你的变胞手来说，这些案例说明：仿人掌弓不仅可以由软体气腔实现，也可以被抽象成特定方向的掌部主动变形自由度。",
    )

    add_heading(doc, "VI. Discussion", 2)
    add_para(
        doc,
        "综述总结，软体机器人手通常具有柔顺结构和较少执行器；在掌部增加一个或多个驱动自由度，可以以高效、鲁棒且安全的方式提升抓取稳定性和/或灵巧性。从机械角度看，驱动掌有时比继续增加手指自由度更简单，因为掌部有更大空间容纳执行器，且执行器可以远离直接操作区域。近几年主动掌文献数量上升，也反映了这一设计趋势。",
    )
    add_para(
        doc,
        "作者归纳出主动掌的五类主要功能：手指重构、新抓取策略、包络抓持、拇指对掌和手内操作。这些也是人手掌的功能，只是在拟人和非拟人夹爪中实现方式不同：拟人手通常尝试复现人手掌弓的形状和运动；非拟人夹爪更关注功能本身，例如通过移动手指基座改变工作空间和抓取类型。",
    )
    add_para(
        doc,
        "论文特别强调，主动掌常用于获得更包络的抓取，即通过扩展接触斑块提升抓持鲁棒性。在拟人手中，这一功能通常与拇指对掌结合；在非拟人夹爪中，掌部驱动主要用于手指重构，改变夹爪工作空间和抓取类型。大多数主动掌只包含 1-2 个驱动自由度，常见驱动方式是气动或电机。",
    )

    add_heading(doc, "VII. Conclusion and Perspectives", 2)
    add_para(
        doc,
        "作者认为，软体手设计领域长期集中于复制人手指的屈伸能力，因此常把手掌做成被动支架，以空间分布手指并获得相对对向关系。被动掌确实有简化设计、控制和可靠性的优势，也能减少执行器和运动机构数量。然而，综述中的案例表明，在掌部加入主动自由度可以释放欠驱动手指无法实现的新抓取和操作能力。",
    )
    add_para(
        doc,
        "未来方向包括：将主动掌与更强传感结合，形成主动触觉感知和交互感知策略；采用新的软体材料、气动板、网状包络结构、可变摩擦或可变黏附表面；进一步从人手掌得到仿生启发，设计能模拟人手解剖结构或能力的掌部，例如拇指对掌和掌部凹腔自适应。作者同时提醒，非拟人结构也有设计自由度，可能用更少执行器实现复杂环境抓取或手内操作。",
    )


def add_tables_and_analysis(doc):
    add_heading(doc, "综述分类框架与关键结论", 1)
    rows = [
        ("关节型 + 非拟人", "掌部多用于重构手指基座位置/方向，改变抓取类型和工作空间。", "齿轮机构、五杆机构、旋转/平移手指基座、scoop 结构。"),
        ("关节型 + 拟人", "掌部多用于拇指对掌、柔性包络和拟人抓持。", "Pisa/IIT SoftHand 类柔性掌，仿人皮肤/柔性掌与欠驱动手指协同。"),
        ("连续型 + 非拟人", "掌部多用于调节夹爪尺寸、工作空间、接触面积、摩擦和预紧。", "气动弯曲掌、伸缩掌、可变摩擦掌、可变法向力掌。"),
        ("连续型 + 拟人", "掌部直接模拟掌弓、掌心凹陷、拇指对掌和对象包络。", "RBO Hand 2/3，单体气动仿生掌，纵弓/远端横弓/斜弓气腔。"),
    ]
    add_table(doc, ["类型", "主动掌主要功能", "代表性设计思路"], rows, [1.35, 2.65, 2.50], font_size=8.4)

    function_rows = [
        ("手指重构", "改变手指基座的位置或姿态，扩展工作空间，改变抓取类型。", "你的四块掌面可被表述为不仅重构掌面，也改变指根相对方位。"),
        ("新抓取策略", "通过掌部附加结构或环境约束形成传统手指闭合无法实现的抓取。", "可强调前折/后折掌形使手能形成不同包络路径。"),
        ("包络抓持", "增加物体与掌/指接触斑块，提高抓持鲁棒性和稳定性。", "这是你三维包络抓持体积的核心功能依据。"),
        ("拇指对掌", "拟人手中常通过掌部或拇指根部自由度改善拇指与其他指的对向关系。", "若你的手有拇指或侧指，可将掌部折叠与对掌能力联系。"),
        ("手内操作", "通过掌部改变支撑、摩擦或法向力，使物体在手内滑动、翻转或重定位。", "可作为未来工作或实验扩展，不一定放在当前主贡献。"),
    ]
    add_table(doc, ["功能", "综述中的含义", "对你的变胞手启发"], function_rows, [1.25, 2.55, 2.70], font_size=8.4)

    add_heading(doc, "创新点与核心贡献", 1)
    contribution_rows = [
        ("科学问题", "机器人手设计长期重视手指而轻视手掌；主动掌在软体/欠驱动手中究竟有哪些形式、功能和趋势，需要系统归纳。"),
        ("研究方法", "Scopus 文献检索 + 综述文献筛选 + 引文追踪；最终纳入 39 篇带主动掌软体手/夹爪文献，按结构、功能、驱动和材料分类。"),
        ("研究内容", "介绍人手掌的生理功能；提出主动掌综述分类；分析关节型/连续型、拟人/非拟人四类装置；总结主动掌五类功能和未来方向。"),
        ("研究目标", "证明掌部可以是主动参与抓取和操作的功能模块，而非仅是承载手指的被动基座。"),
        ("创新点", "把 soft robotic hands 中的 palm 单独作为综述对象，并将主动掌功能归纳为手指重构、新抓取策略、包络抓持、拇指对掌和手内操作。"),
        ("核心贡献", "提供了一个可直接用于论文 Related Work 的主动掌分类框架，并从综述层面支持“少量掌部自由度可提升抓取稳定性和灵巧性”的设计论断。"),
        ("局限性", "综述截止到 2022 年底左右的文献池；主要讨论软体手，刚性/连杆式变胞掌覆盖不足；没有统一量化比较不同主动掌的抓持性能、接触面积或工作空间。"),
    ]
    add_table(doc, ["维度", "分析"], contribution_rows, [1.35, 5.15], font_size=8.6)


def add_relevance(doc):
    add_heading(doc, "与你的变胞手论文的关联分析", 1)
    add_para(
        doc,
        "这里基于你目前的研究对象：球面 6 杆变胞手掌，设置横向轴线与纵向轴线，将手掌分为四块，并希望用远端横弓、纵弓等人手掌弓作为仿生启发，定义可包络抓持的三维掌部变形能力。",
        italic=True,
        color=GRAY,
    )

    rows = [
        ("可直接复用", "主动掌设计动机", "可直接引用本文 Introduction 和 Discussion：机器人手掌常被做成被动支架，但主动掌能用较少自由度增强抓持和操作能力。"),
        ("可直接复用", "人手掌弓仿生依据", "Section II 明确写到掌部沿 longitudinal、distal transverse、oblique arches 弯曲，并形成包绕物体所需的凹形掌面。"),
        ("可直接复用", "包络抓持论据", "Discussion 明确把 enveloping grasp 解释为 extended contact patches，且与更鲁棒的抓取相关。"),
        ("可直接复用", "写作分类", "可把你的工作放入“anthropomorphic / bio-inspired actuated palm”方向，但强调你采用球面 6 杆刚柔/连杆式变胞机构，而非单纯软体气腔。"),
        ("需要改造", "主动掌评价指标", "综述没有统一指标。你的论文可补上三维包络抓持体积、掌面凹度、接触覆盖度、不同折叠状态下的抓取对象范围等量化指标。"),
        ("不宜直接照搬", "软体手材料与驱动路线", "本文大量案例是气动或软体连续结构；你的机构是球面 6 杆和分块掌面，应借鉴功能分类而不是照搬结构实现。"),
    ]
    add_table(doc, ["迁移类型", "内容", "建议"], rows, [1.05, 1.55, 3.90], font_size=8.2)

    add_heading(doc, "可直接写入你 Introduction 的英文表述", 2)
    add_para(
        doc,
        "The palm in robotic hands has often been treated as a passive base for mounting fingers and housing actuation systems. Recent reviews on soft robotic hands with actuated palms show that adding a small number of active degrees of freedom to the palm can enhance enveloping grasping, thumb opposition, finger reconfiguration, and in-hand manipulation, especially for soft or underactuated hands.",
    )
    add_para(
        doc,
        "Inspired by the human palm, which can bend along the longitudinal, distal transverse, and oblique arches to form a concave cavity around objects, an actively reconfigurable palm provides a functional route to improve contact distribution and grasp adaptability.",
    )
    add_para(
        doc,
        "In contrast to soft pneumatic palms, this work realizes palm reconfiguration through a spherical six-bar metamorphic mechanism, enabling coordinated transverse and longitudinal folding of four palm segments for spatial enveloping grasps.",
    )

    add_heading(doc, "建议放在论文中的引用位置", 2)
    cite_rows = [
        ("人手掌弓与凹形掌面", "引用 Pozzi et al. 2024 的 Section II，同时保留 Sangole 2008 和 Neumann 2010。"),
        ("主动掌设计动机", "引用 Pozzi et al. 2024 的 Introduction 和 Conclusion，说明被动掌局限与主动掌趋势。"),
        ("包络抓持和接触斑块", "引用 Pozzi et al. 2024 Discussion 中关于 extended contact patches and robust grasps 的总结。"),
        ("与现有主动掌区别", "在 Related Work 中引用 Pozzi et al. 的四类分类，指出现有多为软体气动/电机掌，而你的贡献是球面 6 杆变胞掌。"),
    ]
    add_table(doc, ["你的论文位置", "推荐使用方式"], cite_rows, [1.75, 4.75], font_size=8.5)

    add_heading(doc, "建议下一步实验/图表", 2)
    next_rows = [
        ("图", "画一个“人手掌弓 - 变胞手横纵轴线 - 四块掌面折叠”的并列示意图，对应本文 Fig. 2(c) 的掌弓逻辑。"),
        ("指标", "定义 transverse folding angle、longitudinal folding angle、palm concavity depth、enveloping grasp volume。"),
        ("对比", "比较 rigid planar palm、only transverse folding、transverse + longitudinal folding 三种构型。"),
        ("实验", "选球、圆柱、椭球和不规则物体，测量接触区域、成功率、抗扰动能力和可包络尺寸范围。"),
    ]
    add_table(doc, ["类型", "建议"], next_rows, [1.10, 5.40], font_size=8.5)


def add_appendix(doc):
    add_heading(doc, "附录：提取与保真说明", 1)
    add_para(doc, "1. 源 PDF 为 IEEE/ASME Transactions on Mechatronics 2024 年第 29 卷第 2 期论文，PDF 元数据显示 DOI 为 10.1109/TMECH.2023.3328944。")
    add_para(doc, "2. PDF 文本层可抽取，报告中的译读基于本地 PDF 文本；原 PDF 部分连字抽取为乱码，已按上下文恢复。")
    add_para(doc, "3. 本报告没有嵌入原论文图像；如需在论文中使用 Fig. 2 或其他原图，应按 IEEE/开放许可要求确认复用条件。")
    add_para(doc, "4. 本报告对你的课题关联分析基于当前研究背景：横向/纵向轴线、四块变胞掌、球面 6 杆机构和三维包络抓持体积。")


def main():
    doc = Document()
    configure_doc(doc)
    add_cover(doc)
    add_metadata_and_terms(doc)
    doc.add_page_break()
    add_section_reading(doc)
    doc.add_page_break()
    add_tables_and_analysis(doc)
    add_relevance(doc)
    add_appendix(doc)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
