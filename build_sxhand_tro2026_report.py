from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os


OUT = os.environ["OUT_DOCX"]
SRC_PDF = os.environ["SRC_PDF"]
SRC_TXT = os.environ["SRC_TXT"]


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Title"].font.size = Pt(18)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(13)
    doc.styles["Heading 3"].font.size = Pt(11.5)


def add_p(doc: Document, text: str, style: str | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p


def add_table(doc: Document, rows, headers):
    data = [headers] + rows
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            r.font.name = "Microsoft YaHei"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            r.font.size = Pt(9)
            if i == 0:
                r.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9EAF7")
                cell._tc.get_or_add_tcPr().append(shading)
    doc.add_paragraph()
    return table


def main() -> None:
    doc = Document()
    setup_styles(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2026 TRO SX-Hand 文献中文翻译与研读报告")
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Design and Implementation of an Anthropomorphic Robotic Hand With Key Kinematic Properties of the Human Hand")
    r.italic = True
    r.font.size = Pt(12)

    add_p(doc, "生成日期：2026-07-05")
    add_p(doc, "分析对象：仿人灵巧手、人体手部运动协同、欠驱动/耦合柔顺机构、主动可动手掌。")
    add_p(doc, "说明：本文档依据本地 PDF 可抽取文本生成。为保持公式和变量可读性，保留英文变量名、图表编号、缩写和关键术语；参考文献列表未逐条翻译，主要作为来源索引处理。")
    doc.add_page_break()

    add_p(doc, "1. 文献信息", "Heading 1")
    add_table(doc, [
        ["论文题目", "Design and Implementation of an Anthropomorphic Robotic Hand With Key Kinematic Properties of the Human Hand"],
        ["作者", "Dai Chu, Jiaji Ma, Siyuan Chen, Jiarui Zhang, Zhiyi Huang, Jinhao Yang, Chang He, Wenbin Chen, Baiyang Sun, Caihua Xiong"],
        ["期刊/年份", "IEEE Transactions on Robotics, Vol. 42, 2026"],
        ["DOI", "10.1109/TRO.2026.3663977"],
        ["原文页数", "21 页"],
        ["源 PDF", SRC_PDF],
        ["抽取文本", SRC_TXT],
        ["提取情况", "标题、摘要、正文、公式变量、图题和主要实验结果均可抽取；少数 IEEE PDF 断行/连字符已在中文整理中人工修正。"],
    ], ["项目", "内容"])

    add_p(doc, "2. 术语表", "Heading 1")
    add_table(doc, [
        ["Anthropomorphic robotic hand", "仿人机器人手/仿人灵巧手", "强调形态、关节布置和运动能力接近人手，而非仅完成夹持。"],
        ["Kinematic synergy (KS)", "运动学协同", "多个关节在任务中呈现的低维协调运动模式。"],
        ["Sparse kinematic synergy (SKS)", "稀疏运动学协同", "只涉及少量强相关关节的协同模式，便于映射到机械耦合单元。"],
        ["JGKSE", "关节分组-运动协同提取方法", "Joint Grouping-KS Extraction，先按相关性分组，再提取局部协同。"],
        ["CMC joint", "腕掌关节/掌骨基底关节", "本文尤其关注无名指和小指 CMC 的可动性，以及拇指 CMC 的对掌能力。"],
        ["Intrafinger coupling-compliance", "指内耦合-柔顺机制", "在同一手指内实现 PIP/DIP 等关节协同，同时保留受阻后的适应性。"],
        ["Interfinger coupling-compliance", "指间耦合-柔顺机制", "在不同手指之间实现强相关关节的机械协同。"],
        ["ORQ", "整体运动重构质量", "衡量保留的协同能否重构整体手部运动数据。"],
        ["IRQ", "单关节运动重构质量", "衡量某个关节的运动是否被协同模式充分重构。"],
        ["Kapandji test", "Kapandji 拇指对掌测试", "常用于评估拇指与其余手指的对掌协调能力。"],
        ["GRASP taxonomy", "GRASP 抓持分类", "Feix 等提出的 33 种人手抓持类型。"],
        ["SX-Hand", "SX-Hand 原型手", "本文提出的 12 驱动、21 关节仿人机器人手，含柔性拇指和 articulated palm。"],
    ], ["英文术语", "中文译名", "说明"])

    add_p(doc, "3. 中文全文翻译与精读", "Heading 1")
    add_p(doc, "3.1 摘要", "Heading 2")
    add_p(doc, "设计一种驱动器数量较少、同时能够复现人手灵巧运动能力的仿人机器人手，仍然是一个具有挑战性的问题。本文提出一种运动学协同分析方法，用于识别运动独立性较高的关节，并量化运动依赖性较强的关节之间的协同特征。基于这些结果，作者提出了指内耦合-柔顺机制和指间耦合-柔顺机制的设计原则，使关节运动学协同能够在机械结构层面得到体现。进一步地，论文形成了一种在减少驱动器数量的同时保留多样运动功能的仿人机器人手设计原则，并将其实现为一个名为 SX-Hand 的原型：该手具有 12 个驱动器、21 个关节、柔性拇指和 articulated palm。实验结果表明，SX-Hand 在 Kapandji 测试中获得最高分，能够完成 GRASP taxonomy 中的全部 33 种抓持类型，并能执行多种灵巧的在手操作任务。该方法为以较少驱动复现丰富运动功能提供了一条有效的概念和技术路径，对高效仿人机器人手和其他仿生系统的设计具有参考价值。")

    sections = [
        ("3.2 引言", [
            "论文从机器人末端执行器的重要性切入，指出人手是自然界中最灵巧的末端执行器之一，因此长期作为机器人手设计的重要仿生对象。与传统夹爪相比，仿人机器人手拥有更多手指、更丰富的关节和更强的任务适应能力，但也带来了驱动、规划和控制复杂度急剧上升的问题。",
            "作者把现有仿人手大体分为两类：一类是接近全驱动的灵巧手，能够提供高自由度，但需要复杂控制和大量执行器；另一类是欠驱动或协同驱动手，能够减少执行器数量，却容易牺牲某些关键运动功能。本文试图在两者之间建立更系统的设计原则：不是简单减少自由度，而是先分析人手哪些关节需要独立、哪些关节适合耦合，再将这种协同关系机械化。",
            "引言特别强调，已有维度降低或协同分析方法通常给出全局性的低维协同，但并不直接告诉设计者“哪些关节应该机械耦合，哪些关节必须保留独立驱动”。这正是本文提出 JGKSE 方法的原因。",
            "作者给出的三项主要贡献包括：提出能兼顾关节独立性和耦合性的稀疏运动协同分析方法；将人手协同特征转化为指内/指间耦合柔顺机构；实现 12 驱动、21 关节 SX-Hand，并通过 Kapandji、工作空间、力学性能、33 类抓持和在手操作实验进行验证。",
        ]),
        ("3.3 人手运动学协同特征分析", [
            "论文建立了 22 关节人手运动学模型，并将与掌部相关的多个关节纳入分析。作者使用两个数据来源：HUST Dataset 包含 30 名被试执行 Feix 等总结的 33 种典型抓持时的 16 个手部关节运动数据；Palm Dataset 补充了掌部关节，尤其是无名指和小指 CMC 屈伸等信息，并包含抓持、操作和手势任务。",
            "作者根据前期研究认为，掌部关节和指间关节之间具有较强独立性，因此将指间关节数据形成 Dataset A，将掌部关节数据形成 Dataset B，分别用于提取不同关节组的运动学协同。这个处理对你的研究很有启发：掌部自由度不应被看作手指运动的附属量，而应作为一个相对独立的结构/运动层进行建模。",
            "作者定义了两个指标来评价运动协同对手部运动的重构质量。ORQ（Overall Reconstruction Quality）用于衡量保留的协同模式对整体手部运动数据的重构能力；IRQ（Individual Reconstruction Quality）用于衡量某个单独关节的运动是否被充分重构。论文采用类似 R² 的形式定义这两个指标，并以 90% 作为有效重构的经验阈值。",
            "传统 PCA 类协同分析往往产生全局协同，即每个协同向量可能同时涉及大量关节。作者认为这种结果不利于机械设计，因为机械耦合更适合对应局部、稀疏、强相关的关节组。为此，论文提出 JGKSE 方法：先根据关节运动相关性进行分组，再在每个组内提取运动学协同。",
            "作者最终识别出 19 个 SKS，用以生成 22 个手部关节的运动。论文中特别提到 RC F/E 和 LC F/E，即无名指与小指 CMC 屈伸具有强耦合关系，可被表示为一个运动耦合单元。这与 active palm 设计高度相关，因为它从人手运动数据层面支持了“尺侧掌部可动”对仿人运动的重要性。",
        ]),
        ("3.4 仿人机器人手设计原则", [
            "基于人手同一手指内 IP 关节的协同特征，作者提出指内耦合-柔顺机制。其基本思路是通过驱动腱-滑轮结构产生主要屈伸运动，再通过耦合腱-滑轮结构在 PIP 和 DIP 等关节之间实现近似人手的协同关系。",
            "柔顺性的关键作用在于：当某个关节因外部物体阻挡而无法继续运动时，其他远端关节仍可继续适应性运动。这避免了刚性耦合导致的适应性不足，使单个驱动器既能产生人手式协同，又能在接触环境中自动调整姿态。",
            "对于不同手指之间存在强运动依赖的关节，作者提出指间耦合-柔顺机制。典型例子是无名指和小指某些关节之间的耦合，以及与掌部 CMC 运动相关的联动。",
            "作者基于上述设计原则实现了 SX-Hand。该手采用腱-滑轮驱动，具有 12 个驱动器和 21 个关节，重量约 1.92 kg，并包含柔性拇指和可动 articulated palm。论文特别强调，SX-Hand 包含无名指和小指 CMC 可动关节，这类关节在已有机器人手中并不常见。",
            "从机构意义上看，SX-Hand 的手掌重点复现的是无名指/小指侧 CMC 可动性，以及由此带来的掌部姿态变化和对掌能力增强。它并不是围绕远端横弓和纵弓构造的球面连杆机构，但其论证方式为你的工作提供了强支撑：掌部关节可动性可以显著扩大手指工作空间，并改善抓持/操作功能。",
        ]),
        ("3.5 SX-Hand 功能评价实验", [
            "Kapandji 测试：SX-Hand 能够完成全部 Kapandji 测试，最终拇指-手指对位姿态接近自然人手。这说明柔性拇指和可动掌部结构能够有效增强拇指与四指之间的协调对掌能力。",
            "指尖可达工作空间：论文计算 actively controllable fingertip reachable workspace，即仅通过驱动器主动驱动即可到达的指尖空间位置。与代表性欠驱动 X-Hand 相比，SX-Hand 的五指总主动可控指尖工作空间约为其 3.4 倍；与人手相比，SX-Hand 的总体主动可控指尖工作空间体积超过人手模型的 90%。",
            "无名指/小指 CMC 可动性的影响：相比 CMC 固定情况，可动 CMC 使无名指工作空间体积增加约 71%，小指工作空间体积增加约 122%。这是一条非常适合在你的论文中引用的证据：手掌/掌骨基底的可动性不是装饰性自由度，而能显著改变手指可达能力。",
            "力学性能：SX-Hand 的指尖输出力超过 10 N，主动抓持力超过 50 N，并可悬吊 20 kg 负载。作者据此说明，尽管该手采用较少驱动器和耦合柔顺结构，仍具备较强的实际抓持负载能力。",
            "抓持与在手操作：在 GRASP taxonomy 的 33 种抓持测试中，SX-Hand 完成了全部 33 种抓持，超过作者此前 X-Hand 的 30 种。论文还展示了多物体抓持和在手操作，包括拧螺母、安装灯泡、旋转两个球、使用剪刀等任务。",
        ]),
        ("3.6 结论", [
            "论文的总体结论是：通过先分析人手关节运动协同，再将强相关关节转化为机械耦合柔顺单元，可以在减少驱动器数量的同时保留丰富的仿人运动功能。SX-Hand 作为 12 驱动、21 关节原型，展示了这种设计路线的可行性。",
            "作者也承认该研究仍有不足，例如缺少复杂动态任务下的长期稳定性评价，耦合机构的参数优化仍依赖设计经验，感知闭环与控制策略仍可进一步发展。",
        ]),
    ]
    for heading, paragraphs in sections:
        add_p(doc, heading, "Heading 2")
        for para in paragraphs:
            add_p(doc, para)

    add_p(doc, "4. 创新点与核心贡献", "Heading 1")
    add_table(doc, [
        ["科学问题", "如何用较少驱动器复现人手丰富的抓持、对掌和在手操作能力。"],
        ["研究方法", "基于人手运动数据的 JGKSE 稀疏运动协同提取；ORQ/IRQ 重构评价；指内/指间耦合柔顺机构设计；原型实验验证。"],
        ["研究内容", "从 22 关节人手模型和数据集出发，识别独立/强耦合关节，提出机械协同实现方法，并研制 SX-Hand。"],
        ["研究目标", "减少 actuator 数量，同时保留人手关键运动学性质和任务功能。"],
        ["创新点 1", "提出面向机械设计的稀疏运动学协同分析，而不是只做全局 PCA 降维。"],
        ["创新点 2", "将运动协同直接映射为指内和指间耦合柔顺机构。"],
        ["创新点 3", "在 12 驱动、21 关节原型中引入较少见的无名指/小指 CMC 可动结构。"],
        ["核心贡献", "给出了一条“人体数据分析—协同提取—机构综合—功能验证”的完整仿人手设计路径。"],
        ["主要证据", "Kapandji 满分、33 种抓持、总工作空间约为 X-Hand 3.4 倍、超过人手模型 90%、CMC 可动使无名指/小指工作空间增加 71%/122%、20 kg 负载。"],
    ], ["项目", "内容"])

    add_p(doc, "5. 局限性与需要谨慎引用之处", "Heading 1")
    for item in [
        "该文的 active palm 重点是 CMC 关节可动、拇指对掌和尺侧掌部运动，而不是显式复现 distal transverse arch 与 longitudinal arch 的三维掌弓结构。若在你的论文中引用，应避免把它表述为“完整掌弓复现”。",
        "SX-Hand 的掌部设计是以关节运动协同和 CMC 可动性为主线，其机构方案与球面六杆/变胞机构不是同类设计。它更适合作为“可动手掌提升功能”的证据，而不是与你的机构拓扑直接同源的先例。",
        "论文报告了大量功能演示，但在复杂接触、长期耐久、闭环触觉控制、接触力分布等方面仍有可扩展空间。",
        "论文把人手关键运动学性质转化为机构设计，但并未给出可对物体包络抓持体积、掌面变形体积或掌弓几何极限的定义。你此前想做的三维包络体积仍需要另行建模。",
    ]:
        add_bullet(doc, item)

    add_p(doc, "6. 与你的变胞手研究的关联分析", "Heading 1")
    add_p(doc, "我理解你的当前研究对象是：设计一种具有球面六杆变胞机构的五指仿人灵巧手，通过横向轴线与纵向轴线将手掌分成四块，以复现人手远端横弓和纵弓，并希望证明这种可动手掌能够提升包络抓持、姿态适应和指尖操作支撑能力。基于这个问题，该文的参考价值如下。")
    add_p(doc, "6.1 可直接复用", "Heading 2")
    for item in [
        "论证逻辑：先从人手关键运动学性质出发，再推导机器人手设计原则。这与你从人手掌弓出发设计变胞掌的写法高度一致。",
        "实验指标：Kapandji 测试、GRASP taxonomy、主动可控指尖工作空间、CMC/掌部自由度对工作空间的增益，这些都可作为你的实验章节参考。",
        "文献对比角度：该文明确说 movable ring/little CMC joints 在已有设计中较少实现，可用于支撑“掌部自由度仍是仿人手设计中的不足环节”。",
        "数据结论：可动 CMC 使无名指/小指工作空间显著增加，说明掌部可动不是边缘功能，而对整体手功能有明确贡献。",
    ]:
        add_bullet(doc, item)

    add_p(doc, "6.2 需要改造后复用", "Heading 2")
    for item in [
        "JGKSE 方法可以启发你分析变胞掌自由度与五指工作空间/抓持功能之间的关系，但你的变量不是传统单关节 CMC，而是球面六杆机构的构型参数和掌块姿态。",
        "SX-Hand 的 articulated palm 可作为“ulnar CMC / palm hollowing”类代表，但不是“横弓+纵弓机构复现”类。",
        "该文工作空间评价可改造成你的“手掌变胞前后五指可达空间/包络抓持体积”评价。",
    ]:
        add_bullet(doc, item)

    add_p(doc, "6.3 不宜直接套用", "Heading 2")
    for item in [
        "该文没有定义掌弓几何极限，也没有给出手掌对物体包络的三维体积定义。",
        "该文没有使用球面连杆、变胞机构或四掌块结构，因此不能作为你球面六杆机构拓扑的直接先例。",
        "该文关注的是运动协同和少驱动设计，不是结构可重构性、机构奇异性或掌弓几何覆盖能力。",
    ]:
        add_bullet(doc, item)

    add_p(doc, "7. 建议你在论文中如何引用这篇文献", "Heading 1")
    for item in [
        "Recent studies have shown that movable palm joints, especially the ring and little finger CMC joints, can substantially increase fingertip reachable workspace and improve anthropomorphic grasping capability.",
        "However, these works mainly focus on CMC mobility, thumb opposition, or kinematic-synergy-driven joint coupling, rather than explicitly realizing the distal transverse and longitudinal arches through a reconfigurable palm mechanism.",
        "In contrast, our design introduces a spherical six-bar metamorphic palm with transverse and longitudinal axes, providing an explicit mechanical realization of human-inspired palm arches.",
    ]:
        add_bullet(doc, item)

    add_p(doc, "8. 可转化为 Fig. 1(b) 或 related-work 机制图的分类", "Heading 1")
    add_table(doc, [
        ["SX-Hand / Chu 2026", "无名指/小指 CMC 可动 + 柔性拇指 + articulated palm", "CMC/palm-hollowing 类", "支撑“掌部可动显著提升功能”，但不是完整横弓纵弓机构"],
        ["RBO Hand 3 / Puhlmann 2022", "软体气动 palm hollowing", "尺侧掌弯曲类", "可作为 soft active palm 对比"],
        ["Pisa/IIT SoftHand palm / Capsi-Morales 2020", "两条 elastic rolling-contact palmar joints", "双轴凹掌类", "与“轴线布置”对比价值高"],
        ["Wang 2021 human-inspired soft palm", "软体掌实现手指展开、掌弯曲、拇指外展/内收", "软体气腔/手指间距调节类", "用于说明已有工作多为软体连续变形"],
        ["你的变胞手", "球面六杆机构；横向轴线 + 纵向轴线；四掌块", "显式掌弓机构复现类", "突出 novelty：机械式复现远端横弓与纵弓，并扩展三维包络能力"],
    ], ["代表工作", "掌部结构", "建议图中类别", "与你工作的关系"])

    add_p(doc, "9. 抽取说明附录", "Heading 1")
    add_p(doc, "PDF 文本抽取共 21 页，约 11.8 万字符。由于 IEEE 双栏 PDF 的断行和连字符，部分标题如 “IV. Functional Evaluation of the SX-Hand” 在抽取文本中存在空格/断行异常；本报告已根据上下文进行人工整理。图像内容未直接复制到本文档中，仅保留图题和对应分析。补充材料和视频未在本次分析中展开。")
    add_p(doc, "注：本报告面向论文写作和课题设计复用，重点突出可动手掌、关节协同、机构设计和实验评价逻辑；不是逐字排版复刻原论文。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
